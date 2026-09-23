"""Parses the merchant's "CONVERSIONE TAGLIE" workbook into normalized JSON.

The workbook is one sheet of stacked per-brand blocks. Row 3 holds the
headers; a block starts wherever column C (BRAND) is non-empty and runs until
the next one. Columns B..AC are the main matrix; everything from column AD
rightwards is per-brand appendices (Ferragamo's proprietary ladder, Hoka's
M/W combined charts, the kids age bands) with repeated, ambiguous headers —
this parser deliberately ignores those and reports them instead of guessing.

Each block yields one ladder per gender actually present:
  - a column named "<X> W" / "<X> M" belongs to the women's / men's ladder
  - a bare column ("EU", "UK", "US", "CM"...) is shared by every ladder in
    the block, which is how the workbook writes brands whose EU run is the
    same for both genders and only the US value differs (adidas, hoka).
A block with no gendered column at all is emitted as `unisex`.

Usage:  python _scripts/parse_conversione_taglie.py <xlsx> [-o out.json]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter

import openpyxl

HEADER_ROW = 3
FIRST_DATA_ROW = 4
AGE_COL, BRAND_COL = 2, 3
# Last column of the main matrix. Everything beyond is per-brand appendix.
LAST_MAIN_COL = 29

# Workbook header -> canonical field in the app's ConversionMapping.
FIELD_BY_HEADER = {
    "EU": "eu",
    "UK": "uk",
    "US": "us",
    "FR": "fr",
    "JPN": "jp",
    "KOR": "kr",
    "CM": "cm",
    "BR": "br",
    "CHINA": "china",
    "MEXICO": "mexico",
    "BG/CH": "bgch",
}


def norm_cell(v):
    """Excel cell -> trimmed string, or None. Keeps .5 and ⅓/⅔ glyphs."""
    if v is None:
        return None
    if isinstance(v, float):
        # 36.666666666666664 -> "36⅔" is NOT done here: the main matrix
        # stores those as text already. A stray float is a real value.
        v = int(v) if v.is_integer() else round(v, 2)
    s = str(v).strip()
    return s or None


def split_header(h: str):
    """('EU W') -> ('eu', 'women'); ('EU') -> ('eu', None)."""
    h = re.sub(r"\s+", " ", str(h or "").strip())
    gender = None
    m = re.match(r"^(.*?)\s+(W|M)$", h)
    if m:
        h, g = m.group(1), m.group(2)
        gender = "women" if g == "W" else "men"
    return FIELD_BY_HEADER.get(h.upper()), gender


def parse(path: str):
    ws = openpyxl.load_workbook(path, data_only=True)["Foglio1"]

    headers = {}
    for c in range(1, LAST_MAIN_COL + 1):
        field, gender = split_header(ws.cell(HEADER_ROW, c).value)
        if field:
            headers[c] = (field, gender)

    starts = [
        r
        for r in range(FIRST_DATA_ROW, ws.max_row + 1)
        if norm_cell(ws.cell(r, BRAND_COL).value)
    ]

    blocks = []
    for i, r0 in enumerate(starts):
        r1 = (starts[i + 1] - 1) if i + 1 < len(starts) else ws.max_row
        brand = norm_cell(ws.cell(r0, BRAND_COL).value)
        age = (norm_cell(ws.cell(r0, AGE_COL).value) or "ADULT").lower()

        genders = {
            g for (_, g) in (headers[c] for c in headers) if g
        }
        present = set()
        for c, (field, gender) in headers.items():
            if any(
                norm_cell(ws.cell(r, c).value) is not None for r in range(r0, r1 + 1)
            ):
                present.add(gender)
        ladders = sorted(g for g in present if g) or ["unisex"]

        rows_by_gender = {g: [] for g in ladders}
        for r in range(r0, r1 + 1):
            for g in ladders:
                row = {}
                for c, (field, gender) in headers.items():
                    if gender is not None and gender != g:
                        continue
                    val = norm_cell(ws.cell(r, c).value)
                    if val is not None:
                        row[field] = val
                if row:
                    rows_by_gender[g].append(row)

        appendix = [
            c
            for c in range(LAST_MAIN_COL + 1, ws.max_column + 1)
            if any(
                norm_cell(ws.cell(r, c).value) is not None for r in range(r0, r1 + 1)
            )
        ]

        blocks.append(
            {
                "brand": brand,
                "age": age,
                "rows": r0,
                "rowsEnd": r1,
                "ladders": {g: rows_by_gender[g] for g in ladders if rows_by_gender[g]},
                "appendixColumns": appendix,
            }
        )
    return blocks


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--out")
    ap.add_argument("--report", action="store_true")
    args = ap.parse_args()

    blocks = parse(args.xlsx)

    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            json.dump(blocks, fh, ensure_ascii=False, indent=1)

    if args.report:
        empty, appendix_only = [], []
        field_use = Counter()
        print(f"{len(blocks)} blocchi\n")
        print(f"{'BRAND':34s} {'AGE':6s} {'LADDER':8s} {'RIGHE':>6s}  CAMPI")
        for b in blocks:
            if not b["ladders"]:
                (appendix_only if b["appendixColumns"] else empty).append(b)
                continue
            for g, rows in b["ladders"].items():
                fields = sorted({k for r in rows for k in r})
                for f in fields:
                    field_use[f] += 1
                print(
                    f"{b['brand'][:34]:34s} {b['age']:6s} {g:8s} {len(rows):6d}  "
                    + ",".join(fields)
                )
        print("\nCampi usati:", dict(field_use.most_common()))
        if appendix_only:
            print("\nSOLO COLONNE APPENDICE (non importati):")
            for b in appendix_only:
                print(f"  {b['brand']:36s} colonne {b['appendixColumns']}")
        if empty:
            print("\nVUOTI:", [b["brand"] for b in empty])
        with_appendix = [b for b in blocks if b["ladders"] and b["appendixColumns"]]
        if with_appendix:
            print("\nCON APPENDICE AGGIUNTIVA (ignorata):")
            for b in with_appendix:
                print(f"  {b['brand']:36s} colonne {b['appendixColumns']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
